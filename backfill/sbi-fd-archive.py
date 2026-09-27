#!/usr/bin/env python3
"""
SBI retail term-deposit history — downloads SBI's own official XLSX and turns it into an
intermediate JSON of "facts only" (dates, tenure labels, amount-band annotations, rates) for
backfill/sbi-fd-archive.ts to turn into RateCards.

Why Python: the repo may not `npm install xlsx` (would touch package.json), so we read the
workbook with openpyxl (pip) instead. This script does NOT interpret tenure labels, amount
bands or dates — it only extracts the raw strings/numbers exactly as the workbook has them.
All domain parsing (parseTenure/parseAmountBand/parseRate/parseDate) happens in the paired
TypeScript loader, per the project's "never guess" rule.

Source: "Domestic Term Deposit Interest Rate: Historical Data" on sbi.bank.in
  Found via: https://sbi.bank.in/web/interest-rates/interest-rates/deposit-rates
  Coverage (as printed): General-public Retail Domestic Term Deposit, 04.01.2008–present.

Workbook layout (verified by hand against the actual file, 2026-09-27):
  - One flat sheet. Columns come in blocks: a "Duration" (or "...Revised Buckets...") column
    lists the tenure-bucket labels for that block, in whatever rows it uses (rows are NOT
    always contiguous — a block may skip a row another block uses). That label set applies to
    every column after it, until the next "Duration"-like column appears.
  - Each non-label column is one rate revision. Its header (row 1) is either an Excel date
    (older columns) or a string such as "Revised for public w.e.f 15/08/2022 (<Rs2cr)" — the
    trailing "(<Rs2cr)"/"(<Rs1cr)"/"(BELOW 15 LAKHS)" style annotation is SBI's own note of the
    retail/bulk threshold in force for that revision, only present on some columns.
  - A handful of cells are footnotes (scheme notes), not rates. We only ever read a cell for a
    row that the block's Duration column defined as a tenure label, and only keep it if it is
    actually numeric — footnotes never survive that filter.

Usage:
  python3 backfill/sbi-fd-archive.py
Output:
  backfill/intermediate/sbi-fd-rows.json
  /agent/workspace/private/archives/sbi/historical-retail-term-deposit-rates.xlsx (raw download)
"""
import json
import os
import re
import sys
import urllib.request
from datetime import datetime

SOURCE_URL = (
    "https://sbi.bank.in/documents/26242/65574/15122025_Historical+Retail+Term+Deposit+Rates.xlsx"
    "/f21a3209-472b-b629-80e8-e558982d4000?t=1765802246329"
)
FOUND_VIA_URL = "https://sbi.bank.in/web/interest-rates/interest-rates/deposit-rates"
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/126.0 Safari/537.36 FixedReturnTracker/1.0 (+https://github.com/SushantKadam73/fixed-return-tracker)"
)

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARCHIVE_DIR = "/agent/workspace/private/archives/sbi"
ARCHIVE_PATH = os.path.join(ARCHIVE_DIR, "historical-retail-term-deposit-rates.xlsx")
OUT_PATH = os.path.join(REPO_ROOT, "backfill", "intermediate", "sbi-fd-rows.json")

LABEL_COL_RE = re.compile(r"^duration\b|revised\s+buckets", re.IGNORECASE)
TRAILING_PAREN_RE = re.compile(r"\(([^()]*)\)\s*$")


def download(url: str, dest: str) -> None:
    if os.path.exists(dest) and os.path.getsize(dest) > 0:
        print(f"  using cached download: {dest}")
        return
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    print(f"  fetching {url}")
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=45) as resp, open(dest, "wb") as f:
        f.write(resp.read())
    print(f"  saved {os.path.getsize(dest)} bytes -> {dest}")


def is_label_column(header) -> bool:
    return isinstance(header, str) and bool(LABEL_COL_RE.search(header.strip()))


def main() -> None:
    try:
        import openpyxl
    except ImportError:
        print("openpyxl is required: pip install --user openpyxl", file=sys.stderr)
        sys.exit(1)

    download(SOURCE_URL, ARCHIVE_PATH)

    wb = openpyxl.load_workbook(ARCHIVE_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]

    revisions = []
    skipped_columns = []
    label_map: dict[int, str] = {}

    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header is None or (isinstance(header, str) and not header.strip()):
            continue  # trailing blank column

        if is_label_column(header):
            new_map: dict[int, str] = {}
            for row in range(2, ws.max_row + 1):
                v = ws.cell(row=row, column=col).value
                if isinstance(v, str) and v.strip():
                    new_map[row] = v.strip()
            label_map = new_map
            continue

        # Rate/revision column.
        effective_date_iso = None
        effective_date_raw = None
        amount_annotation_raw = None
        if isinstance(header, datetime):
            effective_date_iso = header.strftime("%Y-%m-%d")
        else:
            header_str = str(header).strip()
            effective_date_raw = header_str
            m = TRAILING_PAREN_RE.search(header_str)
            if m:
                amount_annotation_raw = m.group(1).strip()

        rows = []
        for row, label in sorted(label_map.items()):
            v = ws.cell(row=row, column=col).value
            if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 < v < 20:
                rows.append({"tenureLabelRaw": label, "rate": round(float(v), 4)})

        if not rows:
            skipped_columns.append(
                {"column": col, "header": effective_date_raw or effective_date_iso, "reason": "no numeric rate cells found for the active tenure labels"}
            )
            continue

        revisions.append(
            {
                "column": col,
                "effectiveDateIso": effective_date_iso,
                "effectiveDateRaw": effective_date_raw,
                "amountAnnotationRaw": amount_annotation_raw,
                "rows": rows,
            }
        )

    out = {
        "sourceUrl": SOURCE_URL,
        "foundViaUrl": FOUND_VIA_URL,
        "extractedAt": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "revisions": revisions,
        "skippedColumns": skipped_columns,
    }
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w") as f:
        json.dump(out, f, indent=1)
        f.write("\n")

    print(f"Extracted {len(revisions)} revision columns, skipped {len(skipped_columns)} columns.")
    print(f"Wrote {OUT_PATH}")
    if skipped_columns:
        print("Skipped columns:")
        for s in skipped_columns:
            print(f"  col {s['column']}: {s['header']!r} — {s['reason']}")


if __name__ == "__main__":
    main()
